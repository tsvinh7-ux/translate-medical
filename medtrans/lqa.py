"""KIỂM ĐỊNH MẪU BỞI CHUYÊN GIA (LQA) và THỬ NGHIỆM MÙ người – máy.

Mục đích: chứng minh chất lượng bằng con số mà người không tin AI cũng kiểm lại được, với công sức con người tối thiểu.

1) Kiểm định mẫu (acceptance sampling, quy tắc "0 lỗi chấp nhận"):
   Muốn tin cậy 95% rằng tỉ lệ đoạn có lỗi nặng ≤ p, chỉ cần chuyên gia đọc n = ln(0,05)/ln(1−p) đoạn chọn NGẪU NHIÊN
   (p = 3% → 99 đoạn; p = 2% → 149 đoạn) – bất kể cuốn sách dài bao nhiêu. Có lỗi → tính trần tỉ lệ lỗi (Wilson 95%),
   bản sửa của chuyên gia được áp vào bản dịch và vào bộ nhớ dịch (nguồn "người").
2) Thử nghiệm mù: với đoạn đã có bản dịch NGƯỜI (TMX nhập từ sách đã xuất bản), xuất cả bản người và bản máy, xáo trộn,
   ẩn nguồn gốc → chuyên gia chấm → mở khoá so sánh điểm người vs máy.
"""
from __future__ import annotations

import json
import math
import random
import time
from pathlib import Path

from .db import Database

SEV = {"nghiêm trọng": "critical", "lớn": "major", "nhỏ": "minor", "critical": "critical", "major": "major",
       "minor": "minor"}
CATS = ["Sai nghĩa / sai chuyên môn", "Bỏ sót / thêm ý", "Thuật ngữ", "Số liệu / đơn vị", "Văn phong / khó đọc",
        "Chính tả / ngữ pháp", "Định dạng"]
SCHEMA = """CREATE TABLE IF NOT EXISTS lqa (
    id INTEGER PRIMARY KEY AUTOINCREMENT, job_id TEXT, kind TEXT, n INTEGER, defects INTEGER, minor INTEGER,
    upper_percent REAL, max_rate_percent REAL, passed INTEGER, details_json TEXT, created_at REAL)"""


def sample_size(confidence: float = 0.95, max_rate: float = 0.03) -> int:
    return math.ceil(math.log(1 - confidence) / math.log(1 - max_rate))


def _wilson_upper(k: int, n: int, z: float = 1.645) -> float:  # một phía 95%
    if n == 0:
        return 1.0
    p = k / n
    denom = 1 + z * z / n
    return min(1.0, (p + z * z / (2 * n) + z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))) / denom)


def _ensure(db: Database) -> None:
    with db._lock:
        db.conn.execute(SCHEMA)
        db.conn.commit()


def latest(db: Database, job_id: str) -> dict | None:
    _ensure(db)
    rows = db.query("SELECT * FROM lqa WHERE job_id=? AND kind='sample' ORDER BY created_at DESC LIMIT 1", (job_id,))
    return dict(rows[0]) if rows else None


def draw(db: Database, job_id: str, n: int, seed: int = 7) -> list[dict]:
    """Chọn mẫu NGẪU NHIÊN phân tầng theo vai trò (tiêu đề / nội dung / chú thích…) trong các đoạn đã qua cổng."""
    segs = [s for s in db.segments(job_id).values()
            if s["kind"] == "text" and s["tgt"] and s["status"] in ("certified", "done", "tm")]
    rnd = random.Random(seed)
    by_role: dict[str, list] = {}
    for s in segs:
        by_role.setdefault(s["role"], []).append(s)
    if len(segs) <= n:
        out = segs
    else:
        out = []
        for role, items in by_role.items():
            k = max(1, round(n * len(items) / len(segs)))
            out.extend(rnd.sample(items, min(k, len(items))))
        rnd.shuffle(out)
        out = out[:n]
    return sorted(out, key=lambda s: (s["page"], s["seg_id"]))


def _plain(text: str, mapping_free: bool = True) -> str:
    import re
    t = re.sub(r"</?(?:b|i|sup|sub)>", "", text or "").replace("⟦¶⟧", " ")
    return re.sub(r"\s+", " ", t).strip()


def _sheet(wb, rows: list[dict], title: str, blind: bool, full: bool = False):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    ws = wb.active
    ws.title = title
    head = ["STT", "Mã", "Trang", "Bản gốc", "Bản dịch", "Kết quả (Đạt/Lỗi)", "Mức lỗi (nghiêm trọng/lớn/nhỏ)",
            "Loại lỗi", "Bản sửa đề xuất (nếu có)", "Ghi chú"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B5394")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 1):
        ws.append([i, r["code"], r.get("page", ""), r["src"], r["tgt"], "", "", "", "", ""])
    for col, w in zip("ABCDEFGHIJ", (5, 12 if not blind else 8, 6, 60, 60, 14, 18, 22, 50, 24)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    last = len(rows) + 1
    for col, opts in (("F", '"Đạt,Lỗi"'), ("G", '"nghiêm trọng,lớn,nhỏ"'), ("H", '"' + ",".join(CATS) + '"')):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "D2"
    guide = wb.create_sheet("Hướng dẫn")
    for line in ([
        "DUYỆT TOÀN VĂN: đọc lần lượt theo thứ tự sách. Đoạn đồng ý thì KHÔNG cần ghi gì.",
        "Đoạn cần sửa: ghi bản sửa hoàn chỉnh vào cột I; chọn 'Lỗi' + mức lỗi + loại lỗi nếu là lỗi (để thống kê); "
        "góp ý không phải lỗi thì chỉ ghi bản sửa / ghi chú.",
    ] if full else []) + [
        "Đọc từng cặp Bản gốc – Bản dịch. Nếu bản dịch dùng được cho xuất bản mà không cần sửa → chọn 'Đạt'.",
        "Nếu có lỗi → chọn 'Lỗi', mức lỗi và loại lỗi; ghi bản sửa đề xuất (nếu có) vào cột I.",
        "nghiêm trọng: sai nghĩa chuyên môn có thể gây hiểu sai lâm sàng (liều, chỉ định, phủ định, số liệu…).",
        "lớn: sai/thiếu ý, sai thuật ngữ, câu khó hiểu – người đọc chuyên môn sẽ phải dừng lại.",
        "nhỏ: văn phong chưa mượt, chính tả, dấu câu – không làm sai nghĩa.",
        "Không thêm/xoá dòng, không đổi cột Mã. Mã dạng ⟦M1⟧, ⟦3⟧ là công thức/ký hiệu giữ nguyên – "
        "khi viết bản sửa hãy giữ đúng các mã này.",
    ] + (["THỬ NGHIỆM MÙ: một số dòng là bản dịch người, một số là bản dịch máy – đánh giá như nhau, "
          "không cần đoán nguồn."] if blind else []):
        guide.append([line])
    guide.column_dimensions["A"].width = 120


def export_sample(db: Database, job_id: str, path: str | Path, n: int | None = None, confidence: float = 0.95,
                  max_rate: float = 0.03, seed: int = 7) -> dict:
    from openpyxl import Workbook
    n = n or sample_size(confidence, max_rate)
    rows = draw(db, job_id, n, seed)
    data = [{"code": s["seg_id"], "page": (s["page"] or 0) + 1, "src": _plain(s["src"]), "tgt": _plain(s["tgt"])}
            for s in rows]
    wb = Workbook()
    _sheet(wb, data, "Kiểm định mẫu", blind=False)
    path = Path(path)
    wb.save(path)
    meta = {"job_id": job_id, "n": len(data), "confidence": confidence, "max_rate": max_rate,
            "required_n": sample_size(confidence, max_rate)}
    path.with_suffix(".meta.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    return {"file": str(path), **meta}


def export_blind(db: Database, job_id: str, path: str | Path, n: int = 60, seed: int = 11) -> dict:
    """Thử nghiệm mù: đoạn có bản dịch NGƯỜI trong bộ nhớ dịch (nguồn human/import) → xuất người + máy, xáo trộn."""
    from openpyxl import Workbook
    from .glossary import normalize_text
    segs = [s for s in db.segments(job_id).values() if s["kind"] == "text" and s["tgt"]]
    pairs = []
    for s in segs:
        rows = db.query("SELECT tgt FROM tm WHERE src_norm=? AND origin IN ('human','import') LIMIT 1",
                        (normalize_text(_plain(s["src"])),))
        if rows and _plain(rows[0]["tgt"]) and _plain(rows[0]["tgt"]) != _plain(s["tgt"]):
            pairs.append((s, _plain(rows[0]["tgt"])))
    rnd = random.Random(seed)
    rnd.shuffle(pairs)
    pairs = pairs[:n]
    items = []
    for s, human in pairs:
        items.append({"src": _plain(s["src"]), "tgt": human, "origin": "human", "seg": s["seg_id"],
                      "page": (s["page"] or 0) + 1})
        items.append({"src": _plain(s["src"]), "tgt": _plain(s["tgt"]), "origin": "ai", "seg": s["seg_id"],
                      "page": (s["page"] or 0) + 1})
    rnd.shuffle(items)
    key = {}
    for i, it in enumerate(items, 1):
        it["code"] = f"B{i:04d}"
        key[it["code"]] = {"origin": it["origin"], "seg": it["seg"]}
    wb = Workbook()
    _sheet(wb, items, "Thử nghiệm mù", blind=True)
    path = Path(path)
    wb.save(path)
    key_path = path.with_suffix(".key.json")  # GIỮ RIÊNG – không gửi cho chuyên gia chấm
    key_path.write_text(json.dumps({"job_id": job_id, "key": key}, ensure_ascii=False), encoding="utf-8")
    return {"file": str(path), "key": str(key_path), "pairs": len(pairs)}


def _read(path: str | Path) -> list[dict]:
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True).worksheets[0]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[1] is None:
            continue
        res = str(row[5] or "").strip().lower()
        sev = SEV.get(str(row[6] or "").strip().lower(), "major" if res.startswith("lỗi") else None)
        out.append({"code": str(row[1]).strip(), "reviewed": bool(res), "error": res.startswith("lỗi"),
                    "severity": sev if res.startswith("lỗi") else None, "category": str(row[7] or ""),
                    "fix": str(row[8] or "").strip(), "note": str(row[9] or "")})
    return out


def _apply_fixes(db: Database, job_id: str, rows: list[dict], domain: str, source: str, reviewer: str = "") -> int:
    """Áp bản sửa của chuyên gia vào bản dịch + bộ nhớ dịch (nguồn "người") + ghi nhận để HỌC VĂN PHONG."""
    import re
    from .learning import record_edit
    segs = db.segments(job_id)
    fixed = 0
    for r in rows:
        s = segs.get(r["code"])
        if not s or not r["fix"]:
            continue
        codes = set(re.findall(r"⟦(?!¶)[^⟧]+⟧", s["src"]))
        if any(c not in r["fix"] for c in codes):  # bản sửa làm mất mã công thức/ký hiệu → giữ làm góp ý
            issues = s["issues"] + [{"severity": "warning", "code": source,
                                     "message": f"Chuyên gia đề xuất (cần biên tập giữ mã {sorted(codes)}): {r['fix']}"}]
            db.execute("UPDATE segments SET status='needs_review', issues_json=?, updated_at=? WHERE job_id=? AND seg_id=?",
                       (json.dumps(issues, ensure_ascii=False), time.time(), job_id, r["code"]))
            continue
        fix = r["fix"]  # đoạn trải nhiều khối: bản sửa không có mốc ⟦¶⟧ → dàn trang chia theo tỉ lệ độ dài
        keep = [i for i in s["issues"] if not i.get("auto") and i.get("severity") not in ("error", "warning", "style")]
        issues = keep + [{"severity": "info", "code": source,
                          "message": f"Chuyên gia sửa ({r['severity'] or 'góp ý'}){': ' + r['note'] if r['note'] else ''}"}]
        db.execute("UPDATE segments SET tgt=?, status='certified', score=100, issues_json=?, updated_at=? "
                   "WHERE job_id=? AND seg_id=?", (fix, json.dumps(issues, ensure_ascii=False), time.time(), job_id,
                                                  r["code"]))
        db.tm_add(s["src"], fix, domain, doc_name=job_id, role=s["role"], quality=100, approved=True, origin="human")
        record_edit(db, job_id, r["code"], domain, s["src"], s["tgt"], fix, source, reviewer)
        fixed += 1
    return fixed


def export_review(db: Database, job_id: str, path: str | Path, scope: str = "all") -> dict:
    """GÓI DUYỆT CHUYÊN GIA toàn văn (song ngữ, theo đúng thứ tự sách): chuyên gia đọc, sửa trực tiếp cột "Bản sửa"."""
    from openpyxl import Workbook
    segs = sorted((s for s in db.segments(job_id).values() if s["kind"] == "text" and s["tgt"]),
                  key=lambda s: (s["page"] or 0, s["seg_id"]))
    if scope == "pending":
        segs = [s for s in segs if s["status"] not in ("certified", "tm")]
    data = [{"code": s["seg_id"], "page": (s["page"] or 0) + 1, "src": _plain(s["src"]), "tgt": _plain(s["tgt"])}
            for s in segs]
    wb = Workbook()
    _sheet(wb, data, "Duyệt chuyên gia", blind=False, full=True)
    path = Path(path)
    wb.save(path)
    return {"file": str(path), "n": len(data), "scope": scope}


def import_review(db: Database, job_id: str, path: str | Path, domain: str = "general_medicine",
                  reviewer: str = "") -> dict:
    """Nhập gói duyệt toàn văn: áp mọi bản sửa, thống kê lỗi chuyên gia phát hiện (đoạn không ghi gì = đồng ý)."""
    _ensure(db)
    rows = _read(path)
    touched = [r for r in rows if r["reviewed"] or r["fix"]]
    fixed = _apply_fixes(db, job_id, touched, domain, "expert_review", reviewer)
    sev = {k: sum(1 for r in touched if r["severity"] == k) for k in ("critical", "major", "minor")}
    res = {"kind": "review", "n": len(rows), "commented": len(touched), "fixed": fixed, **sev, "reviewer": reviewer}
    db.execute("INSERT INTO lqa (job_id, kind, n, defects, minor, upper_percent, max_rate_percent, passed, "
               "details_json, created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
               (job_id, "review", len(rows), sev["critical"] + sev["major"], sev["minor"], 0, 0, 1,
                json.dumps(res, ensure_ascii=False), time.time()))
    return res


SIGNOFF = """CREATE TABLE IF NOT EXISTS signoff (
    id INTEGER PRIMARY KEY AUTOINCREMENT, job_id TEXT, reviewer TEXT, role TEXT, verdict TEXT, note TEXT,
    created_at REAL)"""


def sign_off(db: Database, job_id: str, reviewer: str, verdict: str, role: str = "Chuyên gia hiệu đính",
             note: str = "") -> dict:
    """Chuyên gia ký duyệt: verdict = approve (đồng ý phát hành) | reject (yêu cầu sửa)."""
    with db._lock:
        db.conn.execute(SIGNOFF)
        db.conn.commit()
    if verdict not in ("approve", "reject") or not reviewer.strip():
        raise ValueError("Cần tên người duyệt và kết luận approve | reject")
    db.execute("INSERT INTO signoff (job_id, reviewer, role, verdict, note, created_at) VALUES (?,?,?,?,?,?)",
               (job_id, reviewer.strip(), role, verdict, note, time.time()))
    return latest_signoff(db, job_id)


def latest_signoff(db: Database, job_id: str) -> dict | None:
    with db._lock:
        db.conn.execute(SIGNOFF)
        db.conn.commit()
    rows = db.query("SELECT * FROM signoff WHERE job_id=? ORDER BY id DESC LIMIT 1", (job_id,))
    return dict(rows[0]) if rows else None


def import_sample(db: Database, job_id: str, path: str | Path, max_rate: float = 0.03,
                  apply_fixes: bool = True, domain: str = "general_medicine") -> dict:
    """Nhập kết quả chuyên gia → tỉ lệ lỗi + trần tin cậy 95% → ĐẠT / KHÔNG ĐẠT; áp bản sửa vào bản dịch & TM."""
    _ensure(db)
    rows = [r for r in _read(path) if r["reviewed"]]
    n = len(rows)
    defects = sum(1 for r in rows if r["severity"] in ("critical", "major"))
    minor = sum(1 for r in rows if r["severity"] == "minor")
    upper = _wilson_upper(defects, n)
    passed = n >= sample_size(0.95, max_rate) and defects == 0 or (n > 0 and upper <= max_rate)
    fixed = _apply_fixes(db, job_id, rows, domain, "lqa") if apply_fixes else 0
    by_cat: dict[str, int] = {}
    for r in rows:
        if r["error"]:
            by_cat[r["category"] or "Khác"] = by_cat.get(r["category"] or "Khác", 0) + 1
    res = {"kind": "sample", "n": n, "defects": defects, "minor": minor,
           "upper_percent": round(100 * upper, 2), "max_rate_percent": round(100 * max_rate, 2),
           "passed": bool(passed), "fixed": fixed, "by_category": by_cat}
    db.execute("INSERT INTO lqa (job_id, kind, n, defects, minor, upper_percent, max_rate_percent, passed, "
               "details_json, created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
               (job_id, "sample", n, defects, minor, res["upper_percent"], res["max_rate_percent"], int(passed),
                json.dumps(res, ensure_ascii=False), time.time()))
    return res


def import_blind(db: Database, path: str | Path, key_path: str | Path) -> dict:
    """Mở khoá thử nghiệm mù: so sánh điểm & tỉ lệ lỗi bản NGƯỜI và bản MÁY trên cùng đoạn gốc."""
    _ensure(db)
    key = json.loads(Path(key_path).read_text(encoding="utf-8"))
    w = {"critical": 25, "major": 5, "minor": 1}
    stats = {o: {"n": 0, "ok": 0, "critical": 0, "major": 0, "minor": 0, "penalty": 0} for o in ("human", "ai")}
    for r in _read(path):
        k = key["key"].get(r["code"])
        if not k or not r["reviewed"]:
            continue
        s = stats[k["origin"]]
        s["n"] += 1
        if r["error"]:
            s[r["severity"] or "major"] += 1
            s["penalty"] += w[r["severity"] or "major"]
        else:
            s["ok"] += 1
    out = {}
    for o, s in stats.items():
        n = s["n"] or 1
        out[o] = {**s, "ok_percent": round(100 * s["ok"] / n, 1),
                  "severe_percent": round(100 * (s["critical"] + s["major"]) / n, 1),
                  "score": round(100 - s["penalty"] / n * 4, 1)}
    res = {"kind": "blind", "human": out["human"], "ai": out["ai"],
           "verdict": ("Máy ngang hoặc tốt hơn người" if out["ai"]["score"] >= out["human"]["score"]
                       else "Người tốt hơn máy")}
    db.execute("INSERT INTO lqa (job_id, kind, n, defects, minor, upper_percent, max_rate_percent, passed, "
               "details_json, created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
               (key.get("job_id", ""), "blind", out["ai"]["n"] + out["human"]["n"],
                out["ai"]["critical"] + out["ai"]["major"], out["ai"]["minor"], 0, 0,
                int(out["ai"]["score"] >= out["human"]["score"]), json.dumps(res, ensure_ascii=False), time.time()))
    return res
