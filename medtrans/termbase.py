"""Termbase dự án: xuất ra Excel để kiểm tra (nội bộ hoặc chuyên gia bên ngoài) → nhập lại → KHOÁ trước khi dịch.

Thuật ngữ đã khoá được ghi vào kho thuật ngữ chuẩn (ưu tiên cao nhất) nên được áp dụng xuyên suốt bản dịch,
và được kiểm tra bắt buộc ở từng đoạn.
"""
from __future__ import annotations

from pathlib import Path

from .db import Database, normalize_text

COLUMNS = [
    ("en", "Thuật ngữ gốc", 34), ("vi", "Đề xuất tiếng Việt", 34), ("vi_final", "TIẾNG VIỆT CHỐT (sửa nếu cần)", 34),
    ("status", "Trạng thái (Duyệt / Loại)", 16), ("freq", "Số lần xuất hiện", 10), ("category", "Loại", 14),
    ("confidence", "Độ tin cậy AI", 10), ("origin", "Nguồn đề xuất", 12), ("variants", "Các phương án khác", 30),
    ("context", "Ngữ cảnh trong tài liệu", 60), ("note", "Ghi chú", 40), ("sources", "Dẫn chứng (URL)", 40),
    ("reviewer_note", "Ý kiến người kiểm tra", 30),
]
ORIGIN_VI = {"glossary": "Kho chuẩn", "termprep": "AI đề xuất", "web": "Tra cứu web", "termbase": "Termbase"}
APPROVE = {"duyệt", "duyet", "ok", "approved", "đồng ý", "x", ""}
REJECT = {"loại", "loai", "reject", "rejected", "bỏ", "bo", "không", "khong"}


def export_xlsx(rows: list[dict], path: str | Path, title: str = "") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Termbase"
    head_fill = PatternFill("solid", fgColor="0B5394")
    edit_fill = PatternFill("solid", fgColor="FFF4CE")
    ws.append([c[1] for c in COLUMNS])
    for i, (_, _, w) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    for r in rows:
        vals = []
        for key, _, _ in COLUMNS:
            v = r.get(key, "")
            if key == "vi_final":
                v = r.get("vi", "")
            elif key == "status":
                v = "Duyệt" if r.get("status") != "rejected" else "Loại"
            elif key == "origin":
                v = ORIGIN_VI.get(v, v)
            elif key == "confidence":
                v = round(float(v or 0), 2)
            vals.append(v)
        ws.append(vals)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[2].fill = edit_fill
        row[3].fill = edit_fill
    dv = DataValidation(type="list", formula1='"Duyệt,Loại"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D2:D{max(2, ws.max_row)}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Hướng dẫn")
    for line in [
        f"Termbase dự án: {title}",
        "1. Kiểm tra cột 'Đề xuất tiếng Việt'. Nếu cần sửa, ghi cách dịch chốt vào cột màu vàng 'TIẾNG VIỆT CHỐT'.",
        "2. Cột 'Trạng thái': Duyệt (dùng) hoặc Loại (không dùng). Để trống = Duyệt.",
        "3. Có thể thêm dòng mới (thuật ngữ còn thiếu): điền 'Thuật ngữ gốc' và 'TIẾNG VIỆT CHỐT'.",
        "4. Lưu file và tải lên lại trong MedTrans → 'Khoá termbase' → bắt đầu dịch chính thức.",
        "Thuật ngữ đã khoá được áp dụng thống nhất trong toàn bộ bản dịch và được kiểm tra bắt buộc ở từng đoạn.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    path = Path(path)
    wb.save(path)
    return path


def read_xlsx(path: str | Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Termbase"] if "Termbase" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(it)]
    label_to_key = {c[1]: c[0] for c in COLUMNS}
    keys = [label_to_key.get(h, h) for h in header]
    out = []
    for vals in it:
        r = {k: ("" if v is None else str(v)).strip() for k, v in zip(keys, vals)}
        if r.get("en"):
            out.append(r)
    return out


def lock(db: Database, job_id: str, rows: list[dict], domain: str, src_lang: str = "en") -> dict:
    """Khoá termbase: thuật ngữ duyệt → kho chuẩn (ghi đè, ưu tiên cao nhất); thuật ngữ loại → bỏ."""
    approved = rejected = 0
    with db._lock:
        db.conn.execute("DELETE FROM job_terms WHERE job_id=?", (job_id,))
        db.conn.commit()
    for r in rows:
        status = (r.get("status") or "").strip().lower()
        vi = (r.get("vi_final") or r.get("vi") or "").strip()
        key = normalize_text(r["en"])
        if status in REJECT or not vi:
            rejected += 1
            st = "rejected"
        else:
            approved += 1
            st = "locked"
            db.upsert_term(r["en"], vi, domain, category=r.get("category", ""),
                           note=("Termbase dự án đã kiểm tra. " + (r.get("reviewer_note") or r.get("note") or "")).strip(),
                           source="termbase", priority=90, status="approved", overwrite=True, src_lang=src_lang)
        try:
            freq = int(float(r.get("freq") or 0))
        except ValueError:
            freq = 0
        db.execute("""INSERT OR REPLACE INTO job_terms (job_id, en_norm, en, vi, freq, category, confidence, variants,
                      context, note, origin, sources, status) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                   (job_id, key, r["en"], vi, freq, r.get("category", ""), 1.0, r.get("variants", ""),
                    r.get("context", ""), r.get("reviewer_note") or r.get("note", ""), "termbase",
                    r.get("sources", ""), st))
    return {"locked": approved, "rejected": rejected}


def export_tbx(rows: list[dict], path: str | Path, src_lang: str = "en") -> Path:
    """Xuất TBX (chuẩn trao đổi thuật ngữ, mở được bằng SDL Trados MultiTerm, memoQ...)."""
    from xml.sax.saxutils import escape
    out = ['<?xml version="1.0" encoding="UTF-8"?>', '<tbx type="TBX-Basic" style="dca" xml:lang="vi" '
           'xmlns="urn:iso:std:iso:30042:ed-2">', "<text><body>"]
    for i, r in enumerate(rows):
        if r.get("status") == "rejected":
            continue
        out.append(f'<conceptEntry id="c{i}"><langSec xml:lang="{src_lang}"><termSec><term>{escape(r["en"])}</term>'
                   f'</termSec></langSec><langSec xml:lang="vi"><termSec><term>{escape(r.get("vi", ""))}</term>'
                   f'</termSec></langSec></conceptEntry>')
    out += ["</body></text>", "</tbx>"]
    Path(path).write_text("\n".join(out), encoding="utf-8")
    return Path(path)
