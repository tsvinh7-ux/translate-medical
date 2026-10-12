"""Phân tích tài liệu trước khi dịch (tương tự "Analyze Files" của SDL Trados).

Thống kê: số trang, dung lượng, trang scan, hình, công thức; số đoạn/từ/ký tự cần dịch và không cần dịch;
phân bổ theo mức khớp bộ nhớ dịch (Lặp lại, 100%, 95–99%, 85–94%, 75–84%, 50–74%, Mới); số từ quy đổi (weighted);
thuật ngữ (kho chuẩn + termbase dự án) và tần suất; các câu lặp nhiều nhất; dự báo chi phí/thời gian theo hồ sơ.
"""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from rapidfuzz import fuzz

from .config import PROFILE_ORDER, Config, apply_profile
from .db import Database, normalize_text, text_hash
from .glossary import TermMatcher
from .languages import get_profile
from .pdf_extract import extract

WORD_RE = re.compile(r"[^\W_]+(?:[-'’][^\W_]+)*")
# (khoá, nhãn, hệ số khối lượng công việc còn lại – dùng để quy đổi chi phí/thời gian)
BANDS = [
    ("repetition", "Lặp lại trong tài liệu", 0.0),
    ("tm100", "Khớp 100% (đã chứng nhận)", 0.0),
    ("tm95", "Khớp mờ 95–99%", 0.6),
    ("tm85", "Khớp mờ 85–94%", 0.8),
    ("tm75", "Khớp mờ 75–84%", 0.9),
    ("tm50", "Khớp mờ 50–74%", 1.0),
    ("new", "Mới", 1.0),
]


def words(text: str) -> int:
    return len(WORD_RE.findall(text))


def _band(score: float) -> str:
    if score >= 95:
        return "tm95"
    if score >= 85:
        return "tm85"
    if score >= 75:
        return "tm75"
    if score >= 50:
        return "tm50"
    return "new"


def analyze(pdf_path: str | Path, cfg: Config, db: Database, domain: str | None = None,
            page_range: tuple[int, int] | None = None, with_profiles: bool = True) -> dict:
    from .scan import is_scanned
    import pymupdf

    pdf_path = Path(pdf_path)
    lang = get_profile(cfg.get_path("translation.source_lang", "en"))
    domain = domain or cfg.get_path("translation.domain", "general_medicine")
    domains = list(dict.fromkeys([domain, "general_medicine"]))
    page_numbers = set(range(page_range[0] - 1, page_range[1])) if page_range else None
    pages = extract(pdf_path, lang=lang, ocr_scans=False, page_numbers=page_numbers)
    doc = pymupdf.open(str(pdf_path))
    n_scan = sum(1 for i, pg in enumerate(doc) if (page_numbers is None or i in page_numbers) and is_scanned(pg))
    n_pages_total = doc.page_count
    doc.close()

    tm_min = cfg.get_path("memory.tm_min_quality", 97)
    seen: Counter = Counter()
    band = {k: {"segments": 0, "words": 0, "chars": 0} for k, _, _ in BANDS}
    keep_words = keep_segments = math_blocks = ref_blocks = 0
    texts: list[str] = []
    for p in pages:
        for b in p.blocks:
            if b.keep:
                keep_words += words(b.plain)
                keep_segments += 1
                math_blocks += b.role == "math"
                ref_blocks += b.role == "reference"
                continue
            t = b.plain
            texts.append(t)
            h = text_hash(t)
            seen[h] += 1
            if seen[h] > 1:
                key = "repetition"
            else:
                hit = db.tm_exact(t, domains, src_lang=lang.code)
                if hit and (hit["approved"] or hit["quality"] >= tm_min):
                    key = "tm100"
                else:
                    best = 0.0
                    if len(t) >= 20:
                        norm = normalize_text(t)
                        for c in db.tm_candidates(t, domains, limit=10, src_lang=lang.code):
                            best = max(best, fuzz.ratio(norm, c["src_norm"]))
                    key = _band(best)
            w = words(t)
            band[key]["segments"] += 1
            band[key]["words"] += w
            band[key]["chars"] += len(t)
    # trang scan chưa OCR khi phân tích nhanh → ước lượng 450 từ / 3.000 ký tự mỗi trang
    if n_scan:
        band["new"]["words"] += n_scan * 450
        band["new"]["chars"] += n_scan * 3000
        band["new"]["segments"] += n_scan * 8

    tot_words = sum(v["words"] for v in band.values())
    tot_chars = sum(v["chars"] for v in band.values())
    tot_segs = sum(v["segments"] for v in band.values())
    weight = {k: w for k, _, w in BANDS}
    weighted_words = round(sum(band[k]["words"] * weight[k] for k in band))
    weighted_chars = round(sum(band[k]["chars"] * weight[k] for k in band))
    rows = [{"key": k, "label": label, **band[k], "weight": w,
             "percent": round(100 * band[k]["words"] / tot_words, 1) if tot_words else 0} for k, label, w in BANDS]

    # thuật ngữ
    full = "\n".join(texts)
    full_norm = normalize_text(full)
    matcher = TermMatcher(db.terms(domains, src_lang=lang.code))
    gl_terms = matcher.find(full)
    term_freq = [{"en": t["en"], "vi": t["vi"], "freq": full_norm.count(normalize_text(t["en"])),
                  "origin": "Kho chuẩn"} for t in gl_terms]
    from .translator import job_id_for, pdf_hash_of
    job_id = job_id_for(pdf_hash_of(pdf_path), domain, lang.code)
    project_terms = db.job_terms(job_id)
    known = {normalize_text(t["en"]) for t in term_freq}
    for t in project_terms:
        if normalize_text(t["en"]) not in known:
            term_freq.append({"en": t["en"], "vi": t["vi"], "freq": t["freq"],
                              "origin": "Termbase dự án" + (" (đã khoá)" if t["status"] == "locked" else "")})
    term_freq.sort(key=lambda x: -x["freq"])
    term_occ = sum(t["freq"] for t in term_freq)

    reps = Counter(normalize_text(t) for t in texts)
    top_reps = [{"text": k[:120], "count": c} for k, c in reps.most_common(15) if c > 1]

    result = {
        "file": {"name": pdf_path.name, "size_mb": round(pdf_path.stat().st_size / 1e6, 2),
                 "pages_total": n_pages_total, "pages_analyzed": len(pages), "scanned_pages": n_scan,
                 "images": sum(len(p.images) for p in pages), "math_blocks": math_blocks,
                 "reference_blocks": ref_blocks},
        "totals": {"segments": tot_segs, "words": tot_words, "chars": tot_chars,
                   "a4_pages": round(tot_chars / 3000, 1), "non_translatable_words": keep_words,
                   "non_translatable_segments": keep_segments, "weighted_words": weighted_words,
                   "weighted_chars": weighted_chars,
                   "repetition_words": band["repetition"]["words"],
                   "tm_saving_percent": round(100 * (1 - weighted_chars / tot_chars), 1) if tot_chars else 0},
        "bands": rows,
        "terms": {"unique": len(term_freq), "occurrences": term_occ, "top": term_freq[:40],
                  "project_termbase": len(project_terms),
                  "project_locked": sum(1 for t in project_terms if t["status"] == "locked")},
        "top_repetitions": top_reps,
        "job_id": job_id,
    }
    if with_profiles:
        from .translator import estimate_from_counts
        result["profiles"] = []
        for name in PROFILE_ORDER:
            if name not in (cfg.get("profiles") or {}):
                continue
            pc = apply_profile(cfg, name)
            est = estimate_from_counts(len(pages), weighted_chars, n_scan, result["file"]["images"], pc, db)
            prof = cfg["profiles"][name]
            result["profiles"].append({"name": name, "label": prof.get("label", name),
                                       "description": prof.get("description", ""), **est,
                                       "calibration": calibration_for(db, name, domain)})
    return result


def calibration_for(db: Database, profile: str, domain: str) -> dict | None:
    rows = db.query("SELECT * FROM calibration WHERE profile=? AND (domain=? OR domain='general_medicine') "
                    "ORDER BY created_at DESC LIMIT 1", (profile, domain))
    return dict(rows[0]) if rows else None


def export_xlsx(res: dict, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    head = PatternFill("solid", fgColor="0B5394")

    def sheet(ws, header, rows):
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(70, max(len(str(c.value or ""))
                                                                                  for c in col) + 2))

    ws = wb.active
    ws.title = "Tổng quan"
    f, t = res["file"], res["totals"]
    sheet(ws, ["Chỉ số", "Giá trị"], [
        ["Tệp", f["name"]], ["Dung lượng (MB)", f["size_mb"]], ["Số trang", f["pages_analyzed"]],
        ["Trang scan", f["scanned_pages"]], ["Hình ảnh", f["images"]], ["Khối công thức", f["math_blocks"]],
        ["Số đoạn cần dịch", t["segments"]], ["Số từ cần dịch", t["words"]], ["Số ký tự cần dịch", t["chars"]],
        ["Quy đổi trang A4", t["a4_pages"]], ["Từ không cần dịch (tài liệu tham khảo, công thức...)",
                                               t["non_translatable_words"]],
        ["Số từ quy đổi (weighted)", t["weighted_words"]], ["Tiết kiệm nhờ lặp lại/bộ nhớ dịch (%)",
                                                           t["tm_saving_percent"]],
        ["Thuật ngữ (duy nhất)", res["terms"]["unique"]], ["Lượt xuất hiện thuật ngữ", res["terms"]["occurrences"]]])
    sheet(wb.create_sheet("Mức khớp"), ["Mức khớp", "Số đoạn", "Số từ", "Số ký tự", "% từ", "Hệ số"],
          [[b["label"], b["segments"], b["words"], b["chars"], b["percent"], b["weight"]] for b in res["bands"]])
    sheet(wb.create_sheet("Thuật ngữ"), ["Thuật ngữ gốc", "Tiếng Việt", "Tần suất", "Nguồn"],
          [[x["en"], x["vi"], x["freq"], x["origin"]] for x in res["terms"]["top"]])
    sheet(wb.create_sheet("Câu lặp"), ["Nội dung", "Số lần"], [[x["text"], x["count"]] for x in res["top_repetitions"]])
    if res.get("profiles"):
        sheet(wb.create_sheet("Dự báo theo hồ sơ"),
              ["Hồ sơ", "Chi phí thấp (USD)", "Chi phí cao (USD)", "Thời gian thấp", "Thời gian cao", "Hiệu chuẩn"],
              [[p["label"], p["estimated_cost_usd_low"], p["estimated_cost_usd_high"], p["estimated_time_low"],
                p["estimated_time_high"],
                (f"MQM {p['calibration']['avg_mqm']}, đạt chuẩn {p['calibration']['pass_rate']}%"
                 if p.get("calibration") else "chưa hiệu chuẩn")] for p in res["profiles"]])
    path = Path(path)
    wb.save(path)
    return path
