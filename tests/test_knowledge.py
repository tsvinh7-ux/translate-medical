"""Gói chuyên ngành, học văn phong từ bản sửa của chuyên gia, duyệt toàn văn & ký phát hành."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from make_sample import make  # noqa: E402

from medtrans import learning, lqa  # noqa: E402
from medtrans.assurance import evaluate_release  # noqa: E402
from medtrans.config import apply_profile, load_config  # noqa: E402
from medtrans.db import Database  # noqa: E402
from medtrans.domains import chain, prompt_text, resolve, seed_terms  # noqa: E402
from medtrans.llm import MockLLM  # noqa: E402
from medtrans.translator import Translator, job_id_for  # noqa: E402


def test_domain_packs_resolve_and_inherit(tmp_path):
    assert resolve("Nha khoa") == resolve("dental") == "dentistry"
    assert resolve("Gây mê hồi sức") == "anesthesiology"
    assert chain("dentistry") == ["dentistry", "general_medicine"]
    assert chain("nhi khoa") == ["nhi_khoa", "general_medicine"]          # chưa có gói → vẫn chạy, kế thừa chung
    assert "FDI" in prompt_text("dentistry") and "Y học đại cương" in prompt_text("dentistry")
    assert job_id_for("abc", "Nha khoa") == job_id_for("abc", "dentistry")
    db = Database(tmp_path / "k.db")
    assert seed_terms(db) > 30
    t = {r["en"]: r for r in db.terms(["anesthesiology"])}
    assert t["general anesthesia"]["vi"] == "gây mê toàn thân" and t["general anesthesia"]["priority"] == 70


def test_translator_uses_domain_pack_and_learned_rules(tmp_path):
    cfg = load_config(overrides={"memory": {"db_path": str(tmp_path / "t.db")}})
    db = Database(tmp_path / "t.db")
    db.execute("INSERT INTO tm (src, src_norm, src_hash, tgt, domain, origin, quality, approved, created_at) "
               "VALUES (?,?,?,?,?,?,?,?,?)", ("Periodontitis is a chronic inflammatory disease of the supporting "
                                             "tissues of the teeth caused by bacterial plaque accumulation.",
                                             "x", "h1", "Viêm nha chu là bệnh viêm mạn tính của mô nâng đỡ răng do "
                                             "tích tụ mảng bám vi khuẩn.", "dentistry", "import", 100, 1, 0))
    learning._ensure(db)
    db.execute("INSERT INTO style_rules (domain, kind, rule, status, support, created_at) VALUES (?,?,?,?,?,?)",
               ("dentistry", "style", "Không dùng 'bệnh nhân với'", "approved", 5, 0))
    tr = Translator(apply_profile(cfg, "premium"), db=db, llm=MockLLM({}))
    res = tr.run(make(tmp_path / "d.pdf"), out_dir=tmp_path, domain="Nha khoa")
    assert tr.domain == "dentistry" and tr.domains == ["dentistry", "general_medicine"]
    text = tr.system_translate[-1]["text"]
    assert "Nha khoa" in text and "bệnh nhân với" in text and "VĂN PHONG MẪU" in text and "Viêm nha chu" in text
    assert res["release"]["verdict"] in ("hold", "await_expert")


def test_learning_from_expert_edits(tmp_path):
    db = Database(tmp_path / "l.db")
    for i in range(4):
        assert learning.record_edit(db, "j", f"s{i}", "dentistry", "src", "bệnh nhân với sâu răng",
                                    "người bệnh sâu răng", "review")
    assert not learning.record_edit(db, "j", "s9", "dentistry", "src", "giống nhau", " giống  nhau ", "review")
    new = learning.learn(db, MockLLM({}), "dentistry")
    assert len(new) == 1 and learning.rules(db, "dentistry", "proposed")
    rid = learning.rules(db, "dentistry", "proposed")[0]["id"]
    learning.decide(db, rid, True)
    assert "người bệnh X" in learning.rules_prompt(db, ["dentistry", "general_medicine"])
    assert learning.learn(db, MockLLM({}), "dentistry") == []          # không đề xuất trùng
    db.execute("INSERT INTO style_rules (domain, kind, rule, en, vi, status, support, created_at) VALUES "
               "('dentistry','term','crown (phục hình) → mão răng','prosthetic crown','mão răng','proposed',4,0)")
    tid = learning.rules(db, "dentistry", "proposed")[0]["id"]
    learning.decide(db, tid, True)
    assert any(t["vi"] == "mão răng" for t in db.terms(["dentistry"]))


def test_full_expert_review_and_signoff_gate(tmp_path):
    from openpyxl import load_workbook
    cfg = load_config(overrides={"memory": {"db_path": str(tmp_path / "r.db")}})
    db = Database(tmp_path / "r.db")
    tr = Translator(apply_profile(cfg, "premium"), db=db, llm=MockLLM({}))
    res = tr.run(make(tmp_path / "r.pdf"), out_dir=tmp_path, domain="anesthesiology")
    out = lqa.export_review(db, tr.job_id, tmp_path / "rv.xlsx")
    assert out["n"] >= 5
    wb = load_workbook(tmp_path / "rv.xlsx")
    ws = wb.worksheets[0]
    code = ws.cell(2, 2).value
    src = db.segments(tr.job_id)[code]["src"]
    import re
    fix = ("Bản chuyên gia sửa " + " ".join(re.findall(r"⟦(?!¶)[^⟧]+⟧", src))).strip()
    ws.cell(2, 6, "Lỗi"), ws.cell(2, 7, "nhỏ"), ws.cell(2, 9, fix)
    wb.save(tmp_path / "rv_done.xlsx")
    r = lqa.import_review(db, tr.job_id, tmp_path / "rv_done.xlsx", domain="anesthesiology", reviewer="TS. A")
    assert r["fixed"] == 1 and r["minor"] == 1 and r["n"] == out["n"]
    assert db.segments(tr.job_id)[code]["tgt"] == fix
    assert learning.stats(db)[0]["edits"] == 1                          # bản sửa vào dữ liệu học văn phong
    # cổng phát hành: thiếu chữ ký chuyên gia → không phát hành; ký "đồng ý" → tiêu chí đạt
    for u in tr.units:
        u.status, u.mqm, u.issues = "certified", [], []
    pcfg = apply_profile(cfg, "premium")
    rel = evaluate_release(tr.units, res, pcfg)
    assert rel["verdict"] in ("await_expert", "hold")
    so = lqa.sign_off(db, tr.job_id, "TS. A", "approve", note="Đạt")
    rel2 = evaluate_release(tr.units, res, pcfg, signoff=so)
    expert = next(c for c in rel2["checks"] if "Chuyên gia đọc duyệt" in c["name"])
    assert expert["ok"] and "TS. A" in expert["value"]
    rej = lqa.sign_off(db, tr.job_id, "TS. A", "reject")
    assert evaluate_release(tr.units, res, pcfg, signoff=rej)["verdict"] == "hold"
