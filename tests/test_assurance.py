"""Bảo chứng chất lượng: văn phong, đối soát toàn văn với PDF đầu ra, cổng phát hành, kiểm định mẫu, dự phòng đa AI."""
from __future__ import annotations

import sys
from pathlib import Path

import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).parent))
from make_sample import make  # noqa: E402

from medtrans import lqa  # noqa: E402
from medtrans.assurance import evaluate_release, rendered_coverage  # noqa: E402
from medtrans.config import apply_profile, load_config  # noqa: E402
from medtrans.db import Database  # noqa: E402
from medtrans.llm import LLMError, MockLLM, _backend, make_llm  # noqa: E402
from medtrans.llm_local import FailoverLLM  # noqa: E402
from medtrans.style import document_metrics, segment_issues  # noqa: E402
from medtrans.translator import Translator, Unit  # noqa: E402


def test_style_detects_machine_translationese():
    bad = "Thuốc được chuyển hóa bởi gan một cách nhanh chóng và được thải trừ một cách hiệu quả. Nó rất an toàn."
    codes = {i.code for i in segment_issues(bad)}
    assert {"style_duoc_boi", "style_mot_cach", "style_no_dau_cau"} <= codes
    assert all(i.severity == "style" for i in segment_issues(bad))
    assert segment_issues("Thuốc chẹn beta làm giảm tử vong ở người bệnh suy tim phân suất tống máu giảm.") == []
    m = document_metrics([bad])
    assert m["style_flags"] >= 3 and m["style_flags_per_1000_words"] > 0


def test_rendered_coverage_finds_lost_text(tmp_path):
    pdf = tmp_path / "o.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Heart failure with reduced ejection fraction is defined", fontname="helv")
    doc.save(pdf)
    blocks = {"a": "Heart failure with reduced <b>ejection</b> fraction is defined",
              "b": "This paragraph is not on the output page and must be detected"}
    cov = rendered_coverage(pdf, blocks, {"a": 0, "b": 0})
    assert cov["blocks_missing"] == ["b"] and cov["percent"] == 50.0


@pytest.fixture()
def done_job(tmp_path):
    cfg = load_config(overrides={"memory": {"db_path": str(tmp_path / "a.db")}})
    db = Database(tmp_path / "a.db")
    tr = Translator(apply_profile(cfg, "premium"), db=db, llm=MockLLM({}))
    res = tr.run(make(tmp_path / "s.pdf"), out_dir=tmp_path)
    return cfg, db, tr, res, tmp_path


def test_pipeline_produces_release_verdict_and_certificate(done_job):
    cfg, db, tr, res, _ = done_job
    rel = res["release"]
    assert rel["verdict"] in ("release", "hold") and rel["rendered_percent"] == 100.0
    assert Path(res["certificate"]).exists() and "Phiếu kiểm định" in Path(res["certificate"]).read_text("utf-8")
    names = {c["name"] for c in rel["checks"]}
    assert any("PDF đầu ra" in n for n in names) and any("bỏ sót" in n for n in names)


def test_release_blocks_on_omission(done_job):
    cfg, db, tr, res, _ = done_job
    units = tr.units
    victim = next(u for u in units if u.kind == "text")
    victim.mqm = [{"judge": "fidelity", "severity": "critical", "category": "accuracy_omission"}]
    rel = evaluate_release(units, res, apply_profile(cfg, "premium"))
    assert rel["verdict"] == "hold" and victim.id in rel["to_fix"]
    draft = evaluate_release(units, res, cfg, draft=True)
    assert draft["verdict"] == "draft"


def test_lqa_sampling_and_expert_fix(done_job):
    from openpyxl import load_workbook
    cfg, db, tr, res, tmp = done_job
    assert lqa.sample_size(0.95, 0.03) == 99 and lqa.sample_size(0.95, 0.02) == 149
    out = lqa.export_sample(db, tr.job_id, tmp / "m.xlsx")
    assert out["n"] >= 1
    wb = load_workbook(tmp / "m.xlsx")
    ws = wb.worksheets[0]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 6, "Đạt")
    ws.cell(2, 6, "Lỗi")
    ws.cell(2, 7, "lớn")
    ws.cell(2, 8, "Thuật ngữ")
    code = ws.cell(2, 2).value
    src = db.segments(tr.job_id)[code]["src"]
    fix = ("Bản sửa của chuyên gia " + " ".join(c for c in __import__("re").findall(r"⟦(?!¶)[^⟧]+⟧", src))).strip()
    ws.cell(2, 9, fix)
    wb.save(tmp / "m_done.xlsx")
    r = lqa.import_sample(db, tr.job_id, tmp / "m_done.xlsx")
    assert r["defects"] == 1 and not r["passed"] and r["fixed"] == 1
    seg = db.segments(tr.job_id)[code]
    assert seg["tgt"] == fix and seg["status"] == "certified"
    assert db.query("SELECT origin FROM tm WHERE tgt=?", (fix,))[0]["origin"] == "human"
    assert lqa.latest(db, tr.job_id)["defects"] == 1


def test_blind_test_human_vs_ai(done_job):
    from openpyxl import load_workbook
    cfg, db, tr, res, tmp = done_job
    seg = next(s for s in db.segments(tr.job_id).values() if s["kind"] == "text" and "⟦" not in
               s["src"].replace("⟦¶⟧", ""))
    db.tm_add(lqa._plain(seg["src"]), "Bản dịch của dịch giả đã xuất bản", "cardiology", origin="import",
              approved=True, quality=100)
    out = lqa.export_blind(db, tr.job_id, tmp / "b.xlsx")
    assert out["pairs"] >= 1
    wb = load_workbook(tmp / "b.xlsx")
    ws = wb.worksheets[0]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 6, "Đạt")
    wb.save(tmp / "b_done.xlsx")
    res2 = lqa.import_blind(db, tmp / "b_done.xlsx", out["key"])
    assert res2["human"]["n"] == res2["ai"]["n"] >= 1 and res2["ai"]["ok_percent"] == 100.0


class _Down(MockLLM):
    model = "down"

    def json_call(self, **kw):
        raise LLMError("HTTP 529 overloaded")


def test_failover_switches_provider():
    f = FailoverLLM([_Down({}), MockLLM({})])
    payload = {"segments": [{"id": "x", "source": "Heart failure", "text": "Heart failure"}]}
    out = f.json_call(task="translate", system="", user="", schema={}, payload=payload)
    assert out and f.events and f.events[0]["task"] == "translate"
    with pytest.raises(LLMError):
        FailoverLLM([_Down({}), _Down({})]).json_call(task="translate", system="", user="", schema={}, payload={})


def test_providers_config(monkeypatch):
    monkeypatch.delenv("MEDTRANS_PROVIDER", raising=False)
    llm = load_config()["llm"]
    with pytest.raises(ValueError):
        _backend("openai", llm)                                          # chưa điền model → báo rõ
    cfg = dict(llm, provider="local", failover=["openai"],
               providers={"openai": {"base_url": "http://127.0.0.1:9/v1", "model": "gpt-x",
                                     "price_per_mtok": [1, 2]}})
    f = make_llm(cfg)
    assert isinstance(f, FailoverLLM) and f.chain[1].model == "gpt-x" and f.chain[1].trust_env
    f.chain[1].usage.input_tokens = 1_000_000
    assert f.chain[1].usage.cost("gpt-x") == 1.0


def test_release_blocks_unread_scanned_pages(done_job):
    cfg, db, tr, res, _ = done_job
    rel = evaluate_release(tr.units, res, apply_profile(cfg, "premium"), unread_pages=[3, 7])
    c = next(c for c in rel["checks"] if "Trang scan" in c["name"])
    assert not c["ok"] and rel["verdict"] == "hold" and "trang 3" in rel["to_fix"]


def test_doctor_and_api_key_storage(tmp_path, monkeypatch):
    from medtrans import doctor
    monkeypatch.setenv("MEDTRANS_DB", str(tmp_path / "d.db"))
    items = doctor.run(load_config(), call_api=False)
    names = {i["name"] for i in items}
    assert {"Python", "Thư viện", "Tesseract OCR", "Font tiếng Việt", "Ổ đĩa", "Cơ sở dữ liệu"} <= names
    assert all(i["status"] in ("ok", "warn", "fail") for i in items)
    monkeypatch.setattr(doctor, "ROOT", tmp_path)
    (tmp_path / ".env").write_text("ANTHROPIC_API_KEY=old\nOPENAI_API_KEY=x\n", encoding="utf-8")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    doctor.save_api_key(" sk-ant-new ")
    txt = (tmp_path / ".env").read_text(encoding="utf-8")
    assert "ANTHROPIC_API_KEY=sk-ant-new" in txt and "old" not in txt and "OPENAI_API_KEY=x" in txt
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-...")
    monkeypatch.delenv("MEDTRANS_PROVIDER", raising=False)
    assert doctor.check_api_key(load_config())["status"] == "fail"       # khoá mẫu chưa điền → báo thiếu
